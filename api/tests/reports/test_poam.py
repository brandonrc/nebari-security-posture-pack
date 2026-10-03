import csv
import io
from datetime import timedelta

from openpyxl import load_workbook

from conftest import NOW, VULNS, make_snapshot
from posture.reports.poam import EMASS_COLUMNS, GENERIC_COLUMNS, LEGACY_COLUMNS, TOOL_COLUMNS
from posture.reports.registry import generate

N_FINDINGS = sum(len(v[-1]) for v in VULNS)


def _failing_checks(snapshot):
    """Failing checks that become POA&M rows: probes carry no control (operational hygiene, M4)."""
    return {r.check_id for r in snapshot.posture_results if r.status == "fail"} - {"no-liveness-probe",
                                                                                    "no-readiness-probe"}


def _wb(snapshot, opts, **o):
    rep = generate("poam", "xlsx", snapshot, {**opts, **o})
    assert rep.content_type.startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(rep.content))


def test_xlsx_sheets_and_headers(snapshot, opts):
    wb = _wb(snapshot, opts)
    assert wb.sheetnames == ["eMASS", "POA&M", "eMASS (legacy)", "Info"]
    assert [c.value for c in wb["eMASS"][1]] == EMASS_COLUMNS + TOOL_COLUMNS
    assert [c.value for c in wb["POA&M"][1]] == GENERIC_COLUMNS
    assert [c.value for c in wb["eMASS (legacy)"][1]] == LEGACY_COLUMNS
    assert len(EMASS_COLUMNS) == 32 and len(GENERIC_COLUMNS) == 27 + len(TOOL_COLUMNS)
    assert "Mitigations" in EMASS_COLUMNS  # S1: generic header by default (Navy line is a profile)


def test_row_per_image_cve_package(snapshot, opts):
    wb = _wb(snapshot, opts, poamGranularity="finding")
    expected = N_FINDINGS + len(_failing_checks(snapshot))
    for name in ("eMASS", "POA&M", "eMASS (legacy)"):
        assert wb[name].max_row - 1 == expected, name
    ids = [r[0].value for r in wb["POA&M"].iter_rows(min_row=2)]
    assert len(ids) == len(set(ids)), "POAM IDs must be unique"
    assert any(i.startswith("SP-CFG-") for i in ids)


def test_rollup_by_cve(snapshot, opts):
    wb = _wb(snapshot, opts, rollupByCve=True)
    n_cves = len({v[0] for v in VULNS})
    assert wb["eMASS"].max_row - 1 == n_cves + len(_failing_checks(snapshot))
    ws = wb["POA&M"]
    hdr = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index("POAM ID")] == "SP-CVE-2024-24790")
    assets = row[hdr.index("Asset Identifier")]
    assert "postgres" in assets and "coredns" in assets  # rolled up across two images


def test_sla_and_fields(snapshot, opts):
    ws = _wb(snapshot, opts, poamGranularity="finding")["POA&M"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    r = next(x for x in rows if x["Weakness Source Identifier"] == "CVE-2024-6387")
    detected = (NOW - timedelta(days=40)).date()
    assert r["Original Detection Date"].date() == detected
    assert r["Scheduled Completion Date"].date() == detected + timedelta(days=15)  # critical SLA
    assert "OVERDUE" in r["Comments"]
    assert r["Controls"] == "SI-2"
    assert r["Original Risk Rating"] == "High"
    assert r["Vendor Dependency"] == "No"
    nofix = next(x for x in rows if x["Weakness Source Identifier"] == "CVE-2023-52425")
    assert nofix["Vendor Dependency"] == "Yes" and nofix["Vendor Dependent Product Name"] == "libexpat1"
    assert nofix["Controls"] == "SI-2"
    priv = next(x for x in rows if x["POAM ID"] == "SP-CFG-privileged")
    assert priv["Controls"] == "AC-6, CM-7, SC-39, SC-4"  # SC-4 from the SRG rule's CCI-001090 (N1)
    assert "default/Deployment/legacy-proxy" in priv["Asset Identifier"]
    assert "V-233127" in priv["Weakness Source Identifier"] and "CCI-" in priv["Weakness Source Identifier"]
    # rows are sorted most severe first
    assert rows[0]["Original Risk Rating"] == "High"


def test_emass_sheet_values(snapshot, opts):
    ws = _wb(snapshot, opts, poamGranularity="finding")["eMASS"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    r = next(x for x in rows if x["Security Checks"] == "CVE-2024-45490")
    assert r["POA&M Status"] == "Ongoing"
    assert r["POA&M Item ID"] in (None, "")
    assert r["Controls / APs"] == "SI-2"
    assert r["Raw Severity"] == "Very High" and r["Severity"] in (None, "")  # S1: assessed severity is the ISSO's
    assert r["Mitigations"] in (None, "") and "Rebuild" in r["Recommendations"]  # fix text is not a mitigation
    assert r["External UID"].startswith("SP-CVE-2024-45490-")
    assert r["Milestone ID"] == 1 and r["Milestone Status"] == "Pending"
    assert r["POA&M Scheduled Completion Date"] == r["Milestone Scheduled Completion Date"]
    assert "nginx" in r["Devices Affected"]
    assert ws.cell(row=2, column=hdr.index("POA&M Scheduled Completion Date") + 1).number_format == "mm/dd/yyyy"
    cfg = next(x for x in rows if x["External UID"] == "SP-CFG-no-netpol")
    assert cfg["Controls / APs"] == "AC-4"  # primary in-baseline control (M4)
    assert cfg["Security Checks"].startswith("V-") and "no-netpol" not in cfg["Security Checks"]  # S1
    assert cfg["Identification Source"].startswith("Container Platform Security Requirements Guide")


def test_csv_variants(snapshot, opts):
    opts = {**opts, "poamGranularity": "finding"}
    for variant, cols in (("emass", EMASS_COLUMNS + TOOL_COLUMNS), ("generic", GENERIC_COLUMNS),
                          ("emass-legacy", LEGACY_COLUMNS)):
        rep = generate("poam", "csv", snapshot, {**opts, "poamVariant": variant})
        text = rep.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        assert rows[0] == cols
        assert len(rows) - 1 == N_FINDINGS + len(_failing_checks(snapshot))
        assert all(len(r) == len(cols) for r in rows)
    rows = list(csv.DictReader(io.StringIO(generate("poam", "csv", snapshot, opts).content.decode("utf-8-sig"))))
    assert rows[0]["POA&M Scheduled Completion Date"].count("/") == 2  # MM/DD/YYYY


def test_scope_namespace(opts):
    snap = make_snapshot(scope={"kind": "namespace", "name": "dev"})
    ws = _wb(snap, opts)["POA&M"]
    text = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row)
    assert "nginx" not in text and "jupyterhub" in text
    assert "legacy-proxy" not in text


def test_exclude_system_namespaces(snapshot, opts):
    ws = _wb(snapshot, opts, includeSystemNamespaces=False)["POA&M"]
    text = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row)
    assert "coredns" not in text and "kube-system" not in text


def test_poam_remaps_to_the_baseline_and_derives_controls_from_ccis(snapshot, opts):
    """M4 / N1: rows only carry controls of the selected baseline (eMASS rejects others); items left
    without one are dropped and counted; STIG-derived items add the controls of their CCIs."""
    from posture.reports._common import cci_controls, normalize
    from posture.reports.poam import build_items

    assert cci_controls(["CCI-002233", "CCI-001090", "CCI-002385", "CCI-002605", "CCI-001813", "CCI-003992"]) == [
        "AC-6(8)", "SC-4", "SC-5", "SI-2", "CM-5(1)", "CM-14"]
    v = normalize(snapshot, opts)
    items = build_items(v)
    ids = {i.poam_id for i in items}
    assert "SP-CFG-no-liveness-probe" not in ids and "SP-CFG-no-readiness-probe" not in ids  # no control
    for i in items:
        assert i.controls and all(c in i.controls for c in i.controls)
    rar = next(i for i in items if i.poam_id == "SP-CFG-run-as-root")
    assert "AC-6(8)" not in rar.controls and "AC-6(8)" in rar.controls_outside  # in no baseline
    priv = next(i for i in items if i.poam_id == "SP-CFG-privileged")
    assert priv.ccis and set(priv.controls) >= {"AC-6", "CM-7", "SC-39"}
    assert v.poam_stats["droppedTags"]["AC-6(8)"] >= 1 and v.poam_stats["baseline"] == "moderate"
    low = normalize(snapshot, {**opts, "baseline": "low"})
    assert all("AC-4" not in i.controls for i in build_items(low))  # AC-4 is not in LOW
    wb = _wb(snapshot, opts)
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
    assert info["Control set (baseline)"] == "moderate" and "AC-6(8)" in info["Control tags dropped (outside the baseline)"]


def test_default_rollup_by_remediation_unit(snapshot, opts):
    """S1: one item per image repository (the thing you rebuild), CVEs in Security Checks."""
    wb = _wb(snapshot, opts)
    ws = wb["eMASS"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    repos = [r for r in rows if (r["External UID"] or "").startswith("SP-REPO-")]
    images = {i["repository"] for i in snapshot_images(snapshot)}
    assert 0 < len(repos) <= len(images)
    nginx = next(r for r in repos if "library/nginx" in r["Control Vulnerability Description"])
    assert "CVE-2024-6387" in nginx["Security Checks"] and nginx["Controls / APs"] == "SI-2"
    assert len(rows) == len(repos) + len(_failing_checks(snapshot))
    # stable across regeneration
    again = [r[hdr.index("External UID")] for r in _wb(snapshot, opts)["eMASS"].iter_rows(min_row=2, values_only=True)]
    assert again == [r["External UID"] for r in rows]
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
    assert info["Granularity"].startswith("one item per image repository")


def snapshot_images(snapshot):
    return [{"repository": i.repository} for i in snapshot.images]


def test_overdue_items_are_flagged_for_a_milestone_change(snapshot, opts):
    ws = _wb(snapshot, opts, poamGranularity="finding")["eMASS"]
    hdr = [c.value for c in ws[1]]
    r = next(dict(zip(hdr, x)) for x in ws.iter_rows(min_row=2, values_only=True) if x[hdr.index("Security Checks")]
             == "CVE-2024-6387")
    assert r["Milestone Status Comments"].startswith("OVERDUE") and "milestone change" in r["Milestone Status Comments"]
    assert r["Resources Required"].startswith("Default (confirm")


def test_kev_due_date_overrides_the_sla(opts):
    from posture.reports import kev

    kev.set_catalog({"version": "test", "source": "fixture",
                     "entries": {"CVE-2024-6387": ("2024-07-01", "2026-09-20", 1)}})
    try:
        ws = _wb(make_snapshot(), opts, poamGranularity="finding")["eMASS"]
        hdr = [c.value for c in ws[1]]
        rows = [dict(zip(hdr, x)) for x in ws.iter_rows(min_row=2, values_only=True)]
        r = next(x for x in rows if x["Security Checks"] == "CVE-2024-6387")
        assert r["KEV"] == "Yes" and r["KEV Due Date"].date().isoformat() == "2026-09-20"
        assert r["POA&M Scheduled Completion Date"].date().isoformat() == "2026-09-06"  # SLA was earlier
        other = next(x for x in rows if x["Security Checks"] == "CVE-2024-45490")
        assert other["KEV"] == "No" and other["IAVM ID"] in (None, "")
        kev.set_catalog({"version": "test", "source": "fixture",
                         "entries": {"CVE-2024-45490": ("2024-07-01", "2026-09-01", 0)}})
        ws = _wb(make_snapshot(), opts, poamGranularity="finding")["eMASS"]
        r = next(dict(zip(hdr, x)) for x in ws.iter_rows(min_row=2, values_only=True)
                 if x[hdr.index("Security Checks")] == "CVE-2024-45490")
        assert r["POA&M Scheduled Completion Date"].date().isoformat() == "2026-09-01"  # KEV earlier than SLA
    finally:
        kev.set_catalog(None)


def test_office_org_is_required_and_profiles(snapshot, opts):
    import pytest

    from posture.reports.registry import UnsupportedReport

    snap = make_snapshot(system={"name": "grace", "organization": ""})
    with pytest.raises(UnsupportedReport, match="Office/Org"):
        generate("poam", "xlsx", snap, {**opts, "requireOrganization": True})
    info = {r[0].value: r[1].value for r in _wb(snap, opts)["Info"].iter_rows()}
    assert "eMASS rejects" in info["WARNING"]
    navy = [c.value for c in _wb(snapshot, opts, emassProfile="navy")["eMASS"][1]]
    assert "Mitigations (in-house and in conjunction with the Navy CSSP)" in navy
    army = [c.value for c in _wb(snapshot, opts, emassProfile="army")["eMASS"][1]]
    assert "Predisposing Conditions" not in army and "Threat Description" not in army
    with pytest.raises(UnsupportedReport):
        generate("poam", "xlsx", snapshot, {**opts, "emassProfile": "space-force"})
