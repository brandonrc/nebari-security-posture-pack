"""Product / OS STIG results in the reports (DESIGN §14): per-(image, benchmark) checklists and
stig-bundle.zip, POA&M rows, SAR section, OSCAL AR observations."""

from __future__ import annotations

import io
import json
import xml.etree.ElementTree as ET
import zipfile
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from conftest import NOW, make_snapshot
from posture.reports.registry import UnsupportedReport, generate


def _rule(n, result, sev, cci=None, nist=None, title=None):
    return dict(rule_id=f"xccdf_mil.disa.stig_rule_SV-{n}r1_rule", result=result, severity=sev,
                title=title or f"Rule {n}", stig_id=f"V-{n}", vuln_id=f"V-{n}", sv_id=f"SV-{n}r1_rule",
                rule_version=f"UBTU-22-{n:06d}", cci=cci or [], nist=nist or [], fix_text=f"Fix {n}.",
                group_title="SRG-OS-000480-GPOS-00227",
                first_failed_at=NOW - timedelta(days=40) if result == "fail" else None)


STIG = [
    dict(image_id=2, benchmark_key="ssg-ubuntu2204", benchmark_id="xccdf_org.ssgproject.content_benchmark_UBUNTU_22-04",
         title="Guide to the Secure Configuration of Ubuntu 22.04", version="0.1.82", release_info="",
         source="ssg", profile_id="xccdf_org.ssgproject.content_profile_stig", profile_title="DISA STIG for Ubuntu 22.04",
         content_file="ssg-ubuntu2204-ds.xml", status="evaluated",
         counts={"pass": 2, "fail": 2, "notapplicable": 1, "notchecked": 1}, score=33.3, cat1_open=1, cat2_open=1,
         rootfs_fidelity="full", evaluated_at=NOW - timedelta(hours=1),
         rules=[_rule(260001, "fail", "cat1", cci=["CCI-000366"], title="SSH must not permit root"),
                _rule(260002, "fail", "cat2", nist=["AC-6"]), _rule(260003, "pass", "cat1"),
                _rule(260004, "pass", "cat3"), _rule(260005, "notapplicable", "cat2"),
                _rule(260006, "notchecked", "cat2")]),
    dict(image_id=3, benchmark_key="", status="notApplicable", error="no SCAP benchmark applies (os alpine 3.19)"),
]


@pytest.fixture
def snap():
    return make_snapshot(stig_results=STIG)


def _status(root):
    return {[sd.findtext("ATTRIBUTE_DATA") for sd in v.findall("STIG_DATA")
             if sd.findtext("VULN_ATTRIBUTE") == "Vuln_Num"][0]: v.findtext("STATUS") for v in root.iter("VULN")}


def test_bundle_has_kubernetes_and_product_checklists(snap, opts):
    rep = generate("stig-checklist", "zip", snap, opts)
    assert rep.filename.endswith(".zip") and "stig-bundle" in rep.filename and rep.content_type == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(rep.content))
    names = zf.namelist()
    assert sum(n.startswith("kubernetes/") for n in names) == 2
    prod = [n for n in names if n.startswith("products/") and n.endswith(".ckl")]
    assert len(prod) == 1 and "ssg-ubuntu2204" in prod[0] and "nebari-jupyterhub" in prod[0]
    root = ET.fromstring(zf.read(prod[0]))
    assert root.findtext("ASSET/HOST_NAME") == "quay.io/nebari/nebari-jupyterhub:2024.11.1"
    assert root.findtext("ASSET/HOST_FQDN").endswith("@sha256:" + "b2" * 32)
    st = _status(root)
    assert st == {"V-260001": "Open", "V-260002": "Open", "V-260003": "NotAFinding", "V-260004": "NotAFinding",
                  "V-260005": "Not_Applicable", "V-260006": "Not_Reviewed"}
    vuln = next(v for v in root.iter("VULN") if "V-260001" in ET.tostring(v, encoding="unicode"))
    attrs = {sd.findtext("VULN_ATTRIBUTE"): sd.findtext("ATTRIBUTE_DATA") for sd in vuln.findall("STIG_DATA")}
    assert (attrs["Severity"], attrs["Rule_ID"], attrs["Rule_Ver"], attrs["CCI_REF"]) == (
        "high", "SV-260001r1_rule", "UBTU-22-260001", "CCI-000366")
    cklb = json.loads(zf.read(prod[0] + "b"))
    rules = cklb["stigs"][0]["rules"]
    assert {r["group_id"]: r["status"] for r in rules}["V-260006"] == "not_reviewed"
    assert cklb["target_data"]["host_name"].startswith("quay.io/nebari/")
    index = json.loads(zf.read("products/index.json"))
    assert index["checklists"][0]["fail"] == 2 and index["notEvaluated"][0]["status"] == "notApplicable"


def test_single_product_checklist_and_unknown(snap, opts):
    rep = generate("stig-checklist", "ckl", snap, {**opts, "imageId": 2, "benchmarkId": "ssg-ubuntu2204"})
    assert len(_status(ET.fromstring(rep.content))) == 6 and rep.filename.endswith(".ckl")
    rep = generate("stig-checklist", "cklb", snap, {**opts, "imageId": "2"})
    assert json.loads(rep.content)["stigs"][0]["stig_id"].startswith("xccdf_org.ssgproject")
    with pytest.raises(UnsupportedReport):
        generate("stig-checklist", "ckl", snap, {**opts, "imageId": 3})


def test_poam_rows_for_failing_rules(snap, opts):
    wb = load_workbook(io.BytesIO(generate("poam", "xlsx", snap, opts).content))
    hdr = [c.value for c in wb["eMASS"][1]]
    rows = [dict(zip(hdr, [c.value for c in r])) for r in wb["eMASS"].iter_rows(min_row=2)]
    stig = [r for r in rows if str(r["External UID"]).startswith("SP-STIG-")]
    assert len(stig) == 2
    by = {r["Security Checks"].split()[0]: r for r in stig}
    assert by["V-260001"]["Raw Severity"] == "High" and "CCI-000366" in by["V-260001"]["Security Checks"]
    assert by["V-260001"]["Controls / APs"] == "CM-6"  # CCI-000366 -> CM-6
    assert by["V-260002"]["Raw Severity"] == "Moderate" and by["V-260002"]["Controls / APs"] == "AC-6"
    assert "nebari-jupyterhub" in by["V-260001"]["Devices Affected"]
    assert "Ubuntu 22.04" in by["V-260001"]["Identification Source"]
    # SLA from the first failing evaluation (40 days ago, CAT I -> high -> 30 days): overdue
    assert "OVERDUE" in by["V-260001"]["Comments"]


def test_scope_excludes_other_images(opts):
    snap = make_snapshot(stig_results=STIG, scope={"kind": "namespace", "name": "web"})
    rep = generate("stig-checklist", "zip", snap, opts)
    names = zipfile.ZipFile(io.BytesIO(rep.content)).namelist()
    assert not [n for n in names if n.startswith("products/") and n.endswith(".ckl")]


def test_sar_product_section(snap, opts):
    html = generate("sar", "html", snap, opts).content.decode()
    assert "Product STIG results" in html and "Guide to the Secure Configuration of Ubuntu 22.04" in html
    assert "1 notApplicable" in html
    plain = generate("sar", "html", make_snapshot(), opts).content.decode()
    assert "not enabled or has not evaluated" in plain


def test_oscal_ar_observations(snap, opts):
    from test_oscal import _errors, _fix_patterns  # noqa: F401

    doc = json.loads(generate("oscal-ar", "json", snap, opts).content)
    res = doc["assessment-results"]["results"][0]
    obs = [o for o in res["observations"] if o["title"].startswith("STIG failures")]
    assert len(obs) == 1 and "V-260001" in obs[0]["description"]
    img_items = {i["uuid"] for i in res["local-definitions"]["inventory-items"]
                 if any(p.get("value") == "quay.io/nebari/nebari-jupyterhub:2024.11.1" for p in i["props"])}
    assert obs[0]["subjects"][0]["subject-uuid"] in img_items
    assert any(r["title"].startswith("Image hardening") for r in res["risks"])
    assert any(c["title"] == "OpenSCAP" for c in res["local-definitions"]["components"])


def test_oscal_ar_with_stig_validates(snap, opts):
    import jsonschema

    from test_oscal import SCHEMA, _errors, _fix_patterns

    schema = _fix_patterns(json.loads(SCHEMA.read_text()))
    cls = jsonschema.validators.validator_for(schema)
    v = cls(schema, format_checker=cls.FORMAT_CHECKER)
    assert _errors(v, json.loads(generate("oscal-ar", "json", snap, opts).content)) == []
