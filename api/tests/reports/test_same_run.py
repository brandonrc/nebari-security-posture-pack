"""Compliance review M3: SSP, AR, POA&M and SAR come from the same control evidence run, are stamped
with the run and scan ids, and the AR carries the assertion observations."""

import io
import json

from openpyxl import load_workbook

from conftest import NOW, make_snapshot
from posture.reports.registry import generate

RUN = {"id": 77, "scanId": 42, "baseline": "moderate", "startedAt": "2026-10-01T11:50:00Z",
       "finishedAt": "2026-10-01T11:51:00Z"}
RESULTS = [
    {"id": "kc-brute-force-protection", "title": "Keycloak brute-force detection locks accounts", "status": "fail",
     "controls": ["AC-7"], "component": "keycloak", "severity": "high", "detail": "lockout after 30 failures",
     "evidence": {"failureFactor": 30}, "checkedAt": "2026-10-01T11:50:30Z"},
    {"id": "kc-remember-me-disabled", "title": "'Remember me' is disabled", "status": "pass", "controls": ["AC-12"],
     "component": "keycloak", "severity": "low", "detail": "disabled", "evidence": {}, "checkedAt": "2026-10-01T11:50:30Z"},
]
STATUSES = [
    {"control": "AC-7", "family": "AC", "inBaseline": True, "status": "failing", "assertions": ["kc-brute-force-protection"],
     "objectives": [{"id": "ac-7_obj.a", "state": "not-satisfied"}, {"id": "ac-7_obj.b", "state": "not-satisfied"}]},
    {"control": "AC-12", "family": "AC", "inBaseline": True, "status": "passing", "assertions": ["kc-remember-me-disabled"],
     "objectives": [{"id": "ac-12_obj", "state": "satisfied"}]},
]


def snap():
    return make_snapshot(controls_engine={"data": {"run": RUN, "results": RESULTS, "statuses": STATUSES},
                                          "baseline": "moderate"})


def test_ar_has_assertion_observations_and_engine_based_findings():
    ar = json.loads(generate("oscal-ar", "json", snap(), {"now": NOW}).content)["assessment-results"]
    res = ar["results"][0]
    props = {p["name"]: p["value"] for p in res["props"]}
    assert props["scan-id"] == "42" and props["control-evidence-run"] == "77"
    obs = [o for o in res["observations"] if o["title"].startswith("Control assertion")]
    assert {o["title"].split(":")[0] for o in obs} == {"Control assertion kc-brute-force-protection",
                                                       "Control assertion kc-remember-me-disabled"}
    find = {f["target"]["target-id"]: f["target"]["status"]["state"] for f in res["findings"]}
    assert find["ac-7_obj.a"] == "not-satisfied" and find["ac-7_obj.b"] == "not-satisfied"
    assert find["ac-12_obj"] == "satisfied"
    # scan-only evidence never yields a `satisfied` determination
    assert all(state == "not-satisfied" for oid, state in find.items() if oid != "ac-12_obj")


def test_ar_without_engine_never_claims_satisfied():
    ar = json.loads(generate("oscal-ar", "json", make_snapshot(), {"now": NOW}).content)["assessment-results"]
    assert {f["target"]["status"]["state"] for f in ar["results"][0]["findings"]} == {"not-satisfied"}


def test_poam_has_failing_assertion_items_and_run_stamp():
    wb = load_workbook(io.BytesIO(generate("poam", "xlsx", snap(), {"now": NOW}).content))
    ws = wb["POA&M"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    ctl = [r for r in rows if r[0] == "SP-CTL-kc-brute-force-protection"]
    assert ctl and ctl[0][1] == "AC-7"
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
    assert info["Control evidence run"] == 77 and info["Generated from"] == "scan 42 / control evidence run 77"


def test_ssp_from_same_run_marks_poam_tracked_controls_planned():
    doc = json.loads(generate("oscal-ssp", "json", snap(), {"now": NOW}).content)["system-security-plan"]
    props = {p["name"]: p["value"] for p in doc["system-characteristics"]["props"]}
    assert props["control-evidence-run"] == "77" and props["scan-id"] == "42"
    ac7 = next(r for r in doc["control-implementation"]["implemented-requirements"] if r["control-id"] == "ac-7")
    assert {bc["implementation-status"]["state"] for bc in ac7["by-components"]} == {"planned"}


def test_sar_shows_the_same_run():
    html = generate("sar", "html", snap(), {"now": NOW}).content.decode()
    assert "Control evidence run" in html and "#77" in html
