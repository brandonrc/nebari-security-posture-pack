"""Risk acceptances (controlsEngine.exceptions) in the POA&M, SAR, CRM and Kubernetes STIG checklist."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from conftest import NOW, snapshot_data
from posture.reports.models import ReportSnapshot
from posture.reports.registry import generate

EXC = {"kind": "StatefulSet", "namespace": "dev", "name": "hub-db", "checks": ["run-as-root"],
       "assertions": [], "reason": "vendor image needs root until 2027", "approvedBy": "AO Smith",
       "expiresAt": "2027-01-31", "reviewBy": None, "ticket": "RISK-7", "source": "settings"}
EXPIRED = {**EXC, "kind": "Deployment", "name": "frontend", "namespace": "web", "expiresAt": "2026-01-01", "reason": "old waiver"}


def _snap(exceptions):
    data = snapshot_data()
    for r in data["posture_results"]:
        if (r["namespace"], r["name"], r["check_id"]) == ("dev", "hub-db", "run-as-root"):
            r["status"] = "accepted-risk"
            r["detail"] = "runAsNonRoot not set and runAsUser unset; accepted risk: vendor image needs root until 2027"
    data["controls_engine"] = {"data": {"run": {"id": 9}, "results": [], "statuses": []},
                               "exceptions": exceptions}
    return ReportSnapshot.model_validate(data)


def _sheet(rep, name):
    ws = load_workbook(io.BytesIO(rep.content))[name]
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


def test_poam_risk_acceptance_row():
    rep = generate("poam", "xlsx", _snap([EXC, EXPIRED]), {"now": NOW})
    head, rows = _sheet(rep, "eMASS")
    col = {h: i for i, h in enumerate(head)}
    risk = [r for r in rows if str(r[col["External UID"]]).startswith("SP-RA-")]
    assert len(risk) == 1, "only the active, matching acceptance"
    r = risk[0]
    assert r[col["POA&M Status"]] == "Risk Accepted"
    assert "vendor image needs root until 2027" in r[col["Control Vulnerability Description"]]
    assert r[col["POA&M Requested Risk Accepted Expiration Date"]].date().isoformat() == "2027-01-31"
    assert "AO Smith" in r[col["Comments"]] and "RISK-7" in r[col["Comments"]]
    assert r[col["Mitigations"]] == EXC["reason"]
    # the accepted result is not a failing configuration item any more; the run-as-root item is still
    # there for the other (unaccepted) workloads
    cfg = [x for x in rows if x[col["External UID"]] == "SP-CFG-run-as-root"]
    assert cfg and "dev/StatefulSet/hub-db" not in cfg[0][col["Devices Affected"]]
    head, gen = _sheet(rep, "POA&M")
    g = {h: i for i, h in enumerate(head)}
    gr = [x for x in gen if str(x[g["POAM ID"]]).startswith("SP-RA-")][0]
    assert gr[g["Operational Requirement"]] == "Yes" and gr[g["Deviation Rationale"]] == EXC["reason"]


def test_sar_lists_risk_acceptances():
    html = generate("sar", "html", _snap([EXC, EXPIRED]), {"now": NOW}).content.decode()
    assert "Risk acceptances" in html and "vendor image needs root until 2027" in html
    assert "expires 2027-01-31" in html and "AO Smith" in html
    assert "expired (" in html  # the lapsed waiver is shown as expired


def test_crm_risk_sheet():
    rep = generate("crm", "xlsx", _snap([EXC]), {"now": NOW})
    head, rows = _sheet(rep, "Risk acceptances")
    assert head[0] == "Workload" and rows[0][0] == "dev/StatefulSet/hub-db"
    assert rows[0][3] == EXC["reason"] and rows[0][4] == "AO Smith" and rows[0][8] == "active"
    _, crm = _sheet(rep, "CRM")
    assert any("Risk accepted: dev/StatefulSet/hub-db" in (r[11] or "") for r in crm)


def test_no_exceptions_no_risk_rows():
    rep = generate("poam", "xlsx", _snap([]), {"now": NOW})
    head, rows = _sheet(rep, "eMASS")
    assert not [r for r in rows if str(r[-4]).startswith("SP-RA-")]
    assert "Risk acceptances" not in load_workbook(io.BytesIO(generate("crm", "xlsx", _snap([]), {"now": NOW}).content)
                                                   ).sheetnames


def test_kubernetes_ckl_keeps_accepted_risk_open():
    from posture.reports._common import normalize
    from posture.reports.stig import evaluate_rule, rules_for_check

    snap = _snap([EXC])
    for r in snap.posture_results:  # only the accepted workload fails run-as-root
        if r.check_id == "run-as-root" and r.status == "fail":
            r.status = "pass"
    v = normalize(snap, {"now": NOW})
    rules = [r for r in rules_for_check("run-as-root") if r["evaluation"].get("method") == "posture"]
    assert rules
    res = evaluate_rule(rules[0], v)
    assert res.status == "Open" and "accepted risk" in res.details
