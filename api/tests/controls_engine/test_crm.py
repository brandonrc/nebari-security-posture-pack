"""Customer Responsibility Matrix (compliance review M1 / S6)."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from posture.controls_engine.crm import crm_rows, crm_summary
from posture.controls_engine.engine import derive_statuses
from posture.reports.registry import generate

from .test_ssp import SNAPSHOT, engine_data


def _dicts(rows):
    return [{"control": r.control, "family": r.family, "baseline": r.baseline, "inBaseline": r.in_baseline,
             "status": r.status, "components": r.components, "assertions": r.assertions, "detail": r.detail,
             "responsibility": r.responsibility, "provider": r.provider} for r in rows]


def test_crm_rows_cover_the_baseline_with_responsibility_and_customer_text():
    rows = crm_rows(_dicts(derive_statuses([])))
    by = {r["control"]: r for r in rows}
    assert len(rows) == 287
    assert by["AC-7"]["responsibility"] == "provider" and by["AC-7"]["customerResponsibility"] == ""
    assert by["AC-2"]["responsibility"] == "shared" and "AC-2 a-l" in by["AC-2"]["customerResponsibility"]
    assert "Keycloak" in by["AC-2"]["platformProvides"]
    assert by["AT-2"]["responsibility"] == "org" and "Organization-level" in by["AT-2"]["customerResponsibility"]
    assert by["SC-28"]["responsibility"] == "customer"
    assert by["PE-3"]["responsibility"] == "org" and "hosting" in by["PE-3"]["externallyProvided"]
    s = crm_summary(rows)
    assert sum(s.values()) == 287 and s["provider"] > 0 and s["shared"] > s["provider"]


async def test_crm_report_xlsx_and_csv():
    snap = {**SNAPSHOT, "controls_engine": {"data": await engine_data(), "baseline": "moderate"}}
    rep = generate("crm", "xlsx", snap, {})
    assert rep.filename == "lab-crm-scan42-20261003.xlsx"
    wb = load_workbook(io.BytesIO(rep.content))
    ws = wb["CRM"]
    header = [c.value for c in ws[1]]
    assert header[:4] == ["Control", "Title", "Family", "Responsibility"]
    assert ws.max_row == 288
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
    assert info["Control evidence run"] == 3 and "DRAFT" in info["Report"]
    rep = generate("crm", "csv", SNAPSHOT, {"baseline": "low"})  # no engine data: declarations only
    rows = list(csv.reader(io.StringIO(rep.content.decode("utf-8-sig"))))
    assert len(rows) == 150 and rows[0][0] == "Control"
