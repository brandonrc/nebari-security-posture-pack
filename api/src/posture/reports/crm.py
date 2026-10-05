"""Customer Responsibility Matrix report (`crm`, xlsx / csv): compliance review M1 / S6.

Draft CRM for programs that run on the platform: per control of the selected baseline, the
responsibility (provided / shared / customer / organization), what the platform provides, the
program's residual responsibility and the platform's current evidence status. Built from the
control evidence engine run attached to the snapshot (`controls_engine`)."""

from __future__ import annotations

import csv
import io
from typing import Any

from ._common import TOOL_NAME, filename, get, normalize
from .registry import GeneratedReport

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CT = "text/csv; charset=utf-8"
COLUMNS = ["Control", "Title", "Family", "Responsibility", "Platform evidence status", "Platform provides",
           "Provided outside the platform", "Customer (program) responsibility", "Common control provider",
           "Evidence (assertions)", "Objectives with evidence", "Detail"]


def rows(snapshot: Any, options: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from ..controls_engine.catalog import get_catalog
    from ..controls_engine.crm import crm_rows
    from ..controls_engine.engine import derive_statuses

    ce = get(snapshot, "controls_engine") or {}
    data = ce.get("data") or {}
    baseline = (options.get("baseline") or ce.get("baseline") or "moderate").lower()
    statuses = [dict(s) for s in data.get("statuses") or []]
    if not statuses:  # engine never ran: declarations only
        statuses = [{"control": r.control, "family": r.family, "baseline": r.baseline, "inBaseline": r.in_baseline,
                     "status": r.status, "components": r.components, "assertions": r.assertions, "detail": r.detail,
                     "responsibility": r.responsibility, "provider": r.provider}
                    for r in derive_statuses([], baseline=baseline,
                                             not_applicable=ce.get("notApplicable") or {},
                                             inherit_organizational=bool(ce.get("inheritOrganizationalControls")),
                                             providers=ce.get("commonControlProviders") or [])]
    cat = get_catalog()
    for s in statuses:
        c = cat.get(s["control"])
        s["inBaseline"] = bool(c and c.in_baseline(baseline))
    return crm_rows(statuses), {"baseline": baseline, "run": data.get("run") or {}}


RISK_COLUMNS = ["Workload", "Covers (checks / assertions)", "Controls", "Reason", "Approved by", "Expires",
                "Review by", "Ticket", "Status", "Accepted results", "Source"]


def _risk_rows(v: Any) -> list[list[Any]]:
    out = []
    for ra in v.risk_acceptances:
        e = ra.exception
        status = "expired (findings failing again)" if not ra.active else (
            "active, review overdue" if ra.review_overdue else "active")
        out.append([e.key, ", ".join([*e.checks, *e.assertions]), ", ".join(ra.controls), e.reason, e.approved_by,
                    e.expires_at.isoformat() if e.expires_at else "no expiry",
                    e.review_by.isoformat() if e.review_by else "", e.ticket, status,
                    len(ra.results) + len(ra.assertions), e.source])
    return out


def _annotate_risk(items: list[dict[str, Any]], v: Any) -> None:
    """Controls whose platform evidence includes an accepted risk say so in their Detail column."""
    notes: dict[str, list[str]] = {}
    for ra in v.risk_acceptances:
        if not ra.active or not (ra.results or ra.assertions):
            continue
        e = ra.exception
        exp = e.expires_at.isoformat() if e.expires_at else "no expiry"
        for c in ra.controls:
            notes.setdefault(c, []).append(f"Risk accepted: {e.key} ({', '.join([*ra.checks, *ra.assertion_ids])}; "
                                           f"{e.reason}; approved by {e.approved_by}; {exp})")
    for r in items:
        if r["control"] in notes:
            r["detail"] = " ".join(filter(None, [r["detail"], *notes[r["control"]]]))


def _objectives(r: dict[str, Any]) -> str:
    objs = r.get("objectives") or []
    if not objs:
        return ""
    ev = [o for o in objs if o.get("state") in ("satisfied", "not-satisfied", "evidenced")]
    return f"{len(ev)} of {len(objs)}"


def _values(r: dict[str, Any]) -> list[Any]:
    return [r["control"], r["title"], r["family"], r["responsibilityLabel"], r["status"] or "",
            r["platformProvides"], r["externallyProvided"], r["customerResponsibility"],
            r["commonControlProvider"] or "", ", ".join(r["assertions"]), _objectives(r), r["detail"]]


def generate(fmt: str, snapshot: Any, options: dict[str, Any]) -> GeneratedReport:
    v = normalize(snapshot, options)
    items, meta = rows(snapshot, options)
    _annotate_risk(items, v)
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\r\n")
        w.writerow(COLUMNS)
        for r in items:
            w.writerow(_values(r))
        return GeneratedReport(buf.getvalue().encode("utf-8-sig"), filename(v, "crm", "csv"), CSV_CT)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "CRM"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E7C4FF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in items:
        ws.append(_values(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    widths = [12, 40, 8, 22, 18, 60, 40, 60, 24, 30, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(items) + 1}"
    risk = _risk_rows(v)
    if risk:
        rs = wb.create_sheet("Risk acceptances")
        rs.append(RISK_COLUMNS)
        for c in rs[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E7C4FF")
        for row in risk:
            rs.append(row)
        for row in rs.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        for i, w in enumerate([40, 40, 16, 60, 24, 12, 12, 16, 22, 10, 10], start=1):
            rs.column_dimensions[get_column_letter(i)].width = w
    info = wb.create_sheet("Info")
    counts: dict[str, int] = {}
    for r in items:
        counts[r["responsibilityLabel"]] = counts.get(r["responsibilityLabel"], 0) + 1
    run = meta["run"]
    for k, val in [("Report", "Customer Responsibility Matrix (DRAFT)"), ("System", v.system.name),
                   ("Organization", v.system.organization), ("Baseline", meta["baseline"]),
                   ("Control evidence run", run.get("id") or "none (declarations only)"),
                   ("Evidence checked at", run.get("finishedAt") or ""), ("Scan ID", v.scan.id),
                   ("Generated by", TOOL_NAME), *[(f"Controls - {k}", n) for k, n in counts.items()],
                   ("Risk acceptances", len(risk)),
                   ("Notes", "Draft CRM. 'Provided' and 'Shared' controls become inheritable only after the platform "
                             "itself is assessed and authorized (its own SSP, SAR and ATO, e.g. as a common control "
                             "provider in eMASS). The platform evidence status is machine-collected evidence for an "
                             "assessor, not an assessment result.")]:
        info.append([k, val])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return GeneratedReport(buf.getvalue(), filename(v, "crm", "xlsx"), XLSX_CT)
