"""Plan of Action & Milestones (POA&M) in eMASS import layouts (xlsx, csv).

Workbook sheets:
  * ``eMASS``          - current eMASS (5.x) POA&M import columns (32, "POA&M Item ID" ...
                         "Resulting Residual Risk after Proposed Mitigations").
  * ``POA&M``          - widely published generic / FedRAMP-style POA&M columns.
  * ``eMASS (legacy)`` - pre-5.x DoD POA&M template columns ("Security Control Number" ...).
  * ``Info``           - system, scan, SLA policy and generation notes.

CSV emits one variant (``options.poamVariant``: ``emass`` (default) | ``generic`` |
``emass-legacy``). Column mappings are documented in docs/REPORTS.md.

Granularity (compliance review S1, ``options.poamGranularity``): ``repository`` (default, one item
per image repository = the remediation unit you rebuild, its CVEs in Security Checks), ``cve`` (one
item per CVE with Devices Affected; ``rollupByCve`` is an alias) or ``finding`` (one item per
image x CVE x package). Every item carries a stable External UID. ``options.emassProfile``
(generic | navy | army | marine-corps) adapts the eMASS sheet header to the component's template.
"""

from __future__ import annotations

import csv
import io
from copy import copy
from datetime import datetime
from types import SimpleNamespace
from typing import Any

from ._common import (CAT_SEVERITY, SEVERITIES, TOOL_NAME, View, as_dt, cci_controls, filename, image_label,
                      in_baseline, mdy, normalize, sev, sev_rank, short_hash)
from .cells import safe_row
from .registry import GeneratedReport

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CT = "text/csv; charset=utf-8"
MAX_CELL = 32000  # Excel hard limit is 32767 characters per cell

EMASS_COLUMNS = [
    "POA&M Item ID", "Control Vulnerability Description", "Controls / APs", "Security Checks",
    "POA&M Status", "POA&M Scheduled Completion Date", "POA&M Requested Risk Accepted Expiration Date",
    "POA&M Completion Date", "Milestone ID", "Milestone Description", "Milestone Status",
    "Milestone Status Comments", "Milestone Scheduled Completion Date", "Milestone Completion Date",
    "Identification Source", "Identification Source Details", "Office/Org", "Resources Required", "Comments",
    "Raw Severity", "Devices Affected", "Mitigations",
    "Predisposing Conditions", "Severity", "Relevance of Threat", "Threat Description", "Likelihood", "Impact",
    "Impact Description", "Residual Risk Level", "Recommendations",
    "Resulting Residual Risk after Proposed Mitigations",
]
# Columns after the eMASS template columns (not part of the template; drop them when pasting into it).
TOOL_COLUMNS = ["External UID", "KEV", "KEV Due Date", "IAVM ID"]
# Per-component eMASS template differences (S1): header text and omitted columns.
EMASS_PROFILES: dict[str, dict[str, Any]] = {
    "generic": {},
    "navy": {"rename": {"Mitigations": "Mitigations (in-house and in conjunction with the Navy CSSP)"}},
    "army": {"omit": {"Predisposing Conditions", "Threat Description",
                      "Resulting Residual Risk after Proposed Mitigations"}},
    "marine-corps": {"omit": {"Resulting Residual Risk after Proposed Mitigations"}},
}
GRANULARITIES = ("repository", "cve", "finding")

GENERIC_COLUMNS = [
    "POAM ID", "Controls", "Weakness Name", "Weakness Description", "Weakness Detector Source",
    "Weakness Source Identifier", "Asset Identifier", "Point of Contact", "Resources Required",
    "Overall Remediation Plan", "Original Detection Date", "Scheduled Completion Date", "Planned Milestones",
    "Milestone Changes", "Status Date", "Vendor Dependency", "Last Vendor Check-in Date",
    "Vendor Dependent Product Name", "Original Risk Rating", "Adjusted Risk Rating", "Risk Adjustment",
    "False Positive", "Operational Requirement", "Deviation Rationale", "Supporting Documents", "Comments",
    "Auto-Approve", *TOOL_COLUMNS,
]

LEGACY_COLUMNS = [
    "Control Vulnerability Description", "Security Control Number (NC/NA controls only)", "Office/Org",
    "Security Checks", "Resources Required", "Scheduled Completion Date", "Milestone with Completion Dates",
    "Milestone Changes", "Source Identifying Vulnerability", "Status", "Comments", "Raw Severity",
    "Devices Affected", "Mitigations", "Predisposing Conditions", "Severity", "Relevance of Threat",
    "Threat Description", "Likelihood", "Impact", "Impact Description", "Residual Risk Level",
    "Recommendations", "Resulting Residual Risk after Proposed Mitigations",
]
RESOURCES_DEFAULT = "Default (confirm with the program): existing O&M staff; no additional funding identified."
RESOURCES_VENDOR = "Default (confirm with the program): vendor fix required."
IAVM_NOTE = ""  # IAVM IDs are not available from the scanners; filled in by the ISSO

EMASS_LEVEL = {"critical": "Very High", "high": "High", "medium": "Moderate", "low": "Low",
               "negligible": "Very Low", "unknown": "Low"}
FEDRAMP_RISK = {"critical": "High", "high": "High", "medium": "Moderate", "low": "Low",
                "negligible": "Low", "unknown": "Low"}
CAT = {"critical": "I", "high": "I", "medium": "II", "low": "III", "negligible": "III", "unknown": "III"}


# --------------------------------------------------------------------------- items
def _cut(s: str, n: int = MAX_CELL) -> str:
    s = s or ""
    return s if len(s) <= n else s[: n - 40] + f"\n... [truncated, {len(s) - n + 40} chars omitted]"


def granularity(v: View) -> str:
    g = str(v.options.get("poamGranularity") or ("cve" if v.options.get("rollupByCve") else "repository")).lower()
    if g not in GRANULARITIES:
        from .registry import UnsupportedReport

        raise UnsupportedReport(f"unknown poamGranularity {g!r} ({', '.join(GRANULARITIES)})")
    return g


def _repo_of(v: View, f: SimpleNamespace) -> str:
    img = v.images_by_id.get(f.image_id)
    if img is None:
        return str(f.image_id)
    return "/".join(p for p in (img.registry, img.repository) if p) or img.ref.split("@")[0].rsplit(":", 1)[0]


def _check_list(fs: list[SimpleNamespace], limit: int = 4000) -> str:
    """Vulnerability IDs for Security Checks, KEV and most severe first, within the cell limit."""
    ids: dict[str, tuple] = {}
    for f in fs:
        k = (not f.kev, -sev_rank(f.severity), f.vuln_id)
        ids[f.vuln_id] = min(ids.get(f.vuln_id, k), k)
    out, size = [], 0
    ordered = sorted(ids, key=lambda i: ids[i])
    for n, vid in enumerate(ordered):
        if size + len(vid) + 2 > limit:
            out.append(f"+{len(ordered) - n} more (see vuln-export)")
            break
        out.append(vid)
        size += len(vid) + 2
    return "; ".join(out)


def _vuln_items(v: View, gran: str) -> list[SimpleNamespace]:
    groups: dict[tuple, list] = {}
    for f in v.open_findings:
        key = {"repository": lambda: ("repo", _repo_of(v, f)), "cve": lambda: ("cve", f.vuln_id),
               "finding": lambda: (f.image_id, f.vuln_id, f.package)}[gran]()
        groups.setdefault(key, []).append(f)

    items = []
    for key, fs in groups.items():
        f0 = max(fs, key=lambda f: (bool(f.kev), sev_rank(f.severity)))
        severity = max((f.severity for f in fs), key=sev_rank)
        first_seen = min(f.first_seen_at for f in fs)
        images = [v.images_by_id.get(f.image_id) for f in fs]
        images = list({id(i): i for i in images if i is not None}.values())
        assets = [image_label(i) for i in images] or [str(f0.image_id)]
        workloads = sorted({w for i in images for w in i.workloads})
        scanners = sorted({s for f in fs for s in f.scanners})
        pkgs = sorted({(f.package, f.installed_version, f.fixed_version or "") for f in fs})
        fixable = all(f.fixable for f in fs)
        some_fix = any(f.fixable for f in fs)
        controls = list(dict.fromkeys(c for f in fs for c in f.controls))
        dues = [v.finding_due(f) for f in fs]
        due = min((d for d in dues if d), default=None)
        overdue = bool(due and due < v.now)
        kev = [f for f in fs if f.kev]
        kev_due = min((f.kev_due for f in kev if f.kev_due), default=None)
        sevc: dict[str, int] = {}
        for f in fs:
            sevc[f.severity] = sevc.get(f.severity, 0) + 1
        n_vulns = len({f.vuln_id for f in fs})
        fixes = sorted({f"{p} >= {fx}" for p, _iv, fx in pkgs if fx})
        if gran == "repository":
            repo = key[1]
            uid = f"SP-REPO-{short_hash(repo, n=12)}"
            name = f"{repo}: {n_vulns} vulnerabilit{'y' if n_vulns == 1 else 'ies'}"
            title = f"Vulnerabilities in image repository {repo}"
            desc = (f"{n_vulns} open vulnerabilit{'y' if n_vulns == 1 else 'ies'} ("
                    + ", ".join(f"{n} {s}" for s, n in sorted(sevc.items(), key=lambda x: -sev_rank(x[0])))
                    + f") in {len(pkgs)} package version(s) of image repository {repo}"
                    + (f"; {len(kev)} listed in the CISA KEV catalog" if kev else "") + ".")
            if some_fix:
                plan = (f"Rebuild {repo} on an updated base image / package set ("
                        + "; ".join(fixes[:40]) + (f"; +{len(fixes) - 40} more" if len(fixes) > 40 else "")
                        + "), redeploy, and verify closure with a rescan."
                        + ("" if fixable else " Packages without a published fix: track the vendor advisories, "
                           "document compensating controls or request risk acceptance."))
            else:
                plan = ("No fixed versions are published for these findings. Track the vendor advisories, evaluate "
                        "compensating controls and request risk acceptance if no fix arrives before the scheduled "
                        "completion date.")
            source_id = _check_list(fs)
        else:
            uid = "SP-" + (f0.vuln_id if gran == "cve" else f"{f0.vuln_id}-{short_hash(*key)}")
            name = f"{f0.vuln_id} ({', '.join(sorted({p for p, _, _ in pkgs}))})"
            title = f0.title or f"{f0.vuln_id} in {f0.package}"
            pkg_txt = "; ".join(f"{p} {iv}" + (f" (fixed in {fx})" if fx else " (no fix available)")
                                for p, iv, fx in pkgs)
            desc = title + (f"\n{f0.description}" if f0.description and f0.description != title else "") \
                + f"\nAffected package(s): {pkg_txt}." + (f"\nReference: {f0.url}" if f0.url else "")
            if fixable:
                plan = ("Rebuild the affected image(s) with " + "; ".join(fixes)
                        + " (or update the upstream base image / chart version), redeploy, and verify closure with a"
                          " rescan.")
            else:
                plan = ("No fixed version is published. Track the vendor advisory, evaluate compensating controls "
                        "(network isolation, least privilege, removal of the package if unused) and request risk "
                        "acceptance if the fix will not arrive before the scheduled completion date.")
            source_id = f0.vuln_id
        agreement = max((f.agreement or 0) for f in fs) if any(f.agreement is not None for f in fs) else None
        cvss = max((f.cvss for f in fs if f.cvss is not None), default=None)
        n_enabled = len([s for s in v.scanners if s.enabled]) or 3
        single = sum(1 for f in fs if len(f.scanners) == 1)
        items.append(SimpleNamespace(
            kind="vulnerability", ccis=[], poam_id=uid, source_id=source_id, security_checks=source_id,
            name=name, title=title, description=desc, controls=controls, severity=severity,
            first_seen=first_seen, due=due, overdue=overdue, kev=bool(kev), kev_due=kev_due,
            assets=assets, workloads=workloads, scanners=scanners,
            detector=v.scanner_label(scanners) or "Trivy/Grype/Clair consensus",
            plan=plan, mitigation="", fixable=fixable,
            vendor_product="" if some_fix else ", ".join(sorted({p for p, _, _ in pkgs}))[:2000],
            agreement=agreement, cvss=cvss, url=f0.url,
            identification_source=f"Vulnerability scan - {TOOL_NAME} ({v.scanner_label(scanners)})",
            comments=" ".join(filter(None, [
                f"[External UID {uid}]",
                f"{len(fs)} finding(s) across {len(images)} image digest(s); scanners: "
                f"{', '.join(scanners) or '-'} of {n_enabled} enabled.",
                f"{single} single-scanner finding(s): validate before remediation." if single else "",
                f"CVSS max {cvss}." if cvss is not None else "",
                (f"CISA KEV: {len(kev)} finding(s), due {kev_due:%Y-%m-%d} (BOD 22-01 due date overrides the "
                 "severity SLA)." if kev and kev_due else ""),
                f"SLA {v.sla_days.get(severity)}d for {severity} from first seen {first_seen:%Y-%m-%d} "
                "(per repository, across digests).",
                "OVERDUE." if overdue else "",
            ])),
            impact="",
        ))
    return items


def _posture_items(v: View) -> list[SimpleNamespace]:
    from .stig import rules_for_check  # local import: stig loads YAML lazily

    groups: dict[str, list] = {}
    for r in v.failed_results:
        groups.setdefault(r.check_id, []).append(r)
    items = []
    for check_id, rs in groups.items():
        c = v.check(check_id)
        severity = max((r.severity for r in rs), key=sev_rank)
        first_seen = min((r.first_seen_at for r in rs if r.first_seen_at),
                         default=v.scan.started_at or v.generated_at)
        assets = sorted({r.key + (f" [{r.container}]" if r.container else "") for r in rs})
        stig = rules_for_check(check_id)
        stig_ids = [s["vulnId"] for s in stig]
        from .stig import evaluate_rule

        open_rules = [r for r in stig if evaluate_rule(r, v).status == "Open"]  # S1: cite CKL-Open rules only
        ccis = sorted({c for s in open_rules for c in s.get("ccis") or []})
        checks_txt = "; ".join(f"{r['vulnId']} ({', '.join(r.get('ccis') or [])})" for r in open_rules) \
            or f"Configuration check {check_id} (no Open STIG rule)"
        controls = list(dict.fromkeys([*c.controls, *cci_controls(ccis)]))  # N1: DISA CCI linkage
        details = sorted({r.detail for r in rs if r.detail})
        due = v.sla_due(severity, first_seen)
        overdue = v.overdue(severity, first_seen)
        desc = f"{c.title}: {c.description}".strip().rstrip(":") if c.description else c.title
        desc += f"\nFailing workloads: {len({r.key for r in rs})}."
        if details:
            desc += "\nDetails: " + "; ".join(details[:20]) + (" ..." if len(details) > 20 else "")
        plan = c.remediation or "Update the workload securityContext / manifest to satisfy the check and redeploy."
        sys_note = " (includes Kubernetes system namespaces)" if any(r.system_namespace for r in rs) else ""
        src_stig = ""
        if stig:
            from .stig import benchmark

            b = benchmark(stig[0]["benchmark"])
            src_stig = f"{b['title']} :: Version {b['version']}, {b['releaseInfo']}"
        items.append(SimpleNamespace(
            kind="posture", poam_id=f"SP-CFG-{check_id}", source_id=", ".join([check_id] + stig_ids),
            security_checks=checks_txt, kev=False, kev_due=None,
            name=f"Configuration: {c.title}", title=c.title, description=desc, controls=controls, ccis=ccis,
            severity=severity, first_seen=first_seen, due=due, overdue=overdue, assets=assets,
            workloads=sorted({r.key for r in rs}), scanners=[], detector=f"{TOOL_NAME} posture check {check_id}",
            plan=plan, mitigation="", fixable=True, vendor_product="", agreement=None, cvss=None, url="",
            identification_source=src_stig or f"Configuration review - {TOOL_NAME}",
            comments=" ".join(filter(None, [
                f"[External UID SP-CFG-{check_id}]",
                f"Kubernetes posture check '{check_id}' failed on {len(rs)} container(s){sys_note}.",
                f"Mapped STIG rules: {', '.join(stig_ids)}." if stig_ids else "",
                f"SLA {v.sla_days.get(severity)}d from first seen {first_seen:%Y-%m-%d}.",
                "OVERDUE." if overdue else "",
            ])),
            impact="",
        ))
    return items


def _stig_items(v: View) -> list[SimpleNamespace]:
    """Failing product / OS STIG rules (DESIGN §14), one item per (image, rule): CAT -> raw severity,
    Security Checks = V-ID (CCIs), controls from the DISA CCI map (else the content's NIST 800-53
    references, else CM-6), SLA clock from the first evaluation that found the rule failing."""
    items = []
    for b, r in v.stig_failures:
        img = b.image
        severity = CAT_SEVERITY.get(r.severity, "low")
        first_seen = r.first_failed_at or b.evaluated_at or v.scan.started_at or v.generated_at
        due, overdue = v.sla_due(severity, first_seen), v.overdue(severity, first_seen)
        vid = r.vuln_id or r.stig_id or r.rule_version or r.rule_id
        controls = cci_controls(r.cci) or list(dict.fromkeys(r.nist or [])) or ["CM-6"]
        uid = f"SP-STIG-{short_hash(img.digest or img.ref, b.benchmark_key, r.rule_id, n=10)}"
        stig_ref = f"{b.title} :: Version {b.version}" + (f", {b.release_info}" if b.release_info else "")
        checks = f"{vid}" + (f" ({', '.join(r.cci)})" if r.cci else "") + (
            f" [{r.rule_version}]" if r.rule_version and r.rule_version != vid else "")
        fid = " Evaluated on a rootfs extracted without root (degraded fidelity)." if b.rootfs_fidelity == "degraded" else ""
        items.append(SimpleNamespace(
            kind="stig", poam_id=uid, source_id=vid, security_checks=checks, kev=False, kev_due=None,
            name=f"Product STIG: {r.title}", title=r.title,
            description=f"{r.title} ({b.title}, profile {b.profile_title or b.profile_id}) fails in container image "
                        f"{image_label(img)}.",
            controls=controls, ccis=list(r.cci or []), severity=severity, first_seen=first_seen, due=due,
            overdue=overdue, assets=[image_label(img)], workloads=list(img.workloads or []), scanners=[],
            detector=f"OpenSCAP ({TOOL_NAME}) {b.benchmark_id or b.benchmark_key} profile {b.profile_id}",
            plan=r.fix_text or f"Remediate {vid} in the image build (Dockerfile / base image) and rebuild.",
            mitigation="", fixable=True, vendor_product=img.ref, agreement=None, cvss=None, url="",
            identification_source=stig_ref, impact="",
            comments=" ".join(filter(None, [
                f"[External UID {uid}]", f"CAT {'I' * {'cat1': 1, 'cat2': 2, 'cat3': 3}.get(r.severity, 3)}.",
                f"Rule {r.rule_id}.", f"SLA {v.sla_days.get(severity)}d from first failure {first_seen:%Y-%m-%d}.",
                "OVERDUE." if overdue else "", fid.strip()])),
        ))
    return items


def _assertion_items(v: View) -> list[SimpleNamespace]:
    """Failing platform control assertions of the attached control evidence run (M3): the same run
    the SSP / AR / SAR use, so every failing control has a POA&M item."""
    run = v.engine_run
    started = as_dt(run.get("startedAt") or run.get("finishedAt")) or v.scan.started_at or v.generated_at
    items = []
    for r in v.engine_results:
        if r.get("status") != "fail":
            continue
        severity = sev(r.get("severity") or "medium")
        first_seen = as_dt(r.get("firstFailedAt")) or started
        due, overdue = v.sla_due(severity, first_seen), v.overdue(severity, first_seen)
        aid = r.get("id", "")
        detail = (r.get("detail") or "").strip()
        items.append(SimpleNamespace(
            kind="assertion", ccis=[], poam_id=f"SP-CTL-{aid}", source_id=aid,
            security_checks=f"Control assertion {aid}", kev=False, kev_due=None, name=f"Platform control: {r.get('title', aid)}",
            title=r.get("title", aid), description=f"{r.get('title', aid)} (assertion {aid}) fails: {detail}",
            controls=list(r.get("controls") or []), severity=severity, first_seen=first_seen, due=due,
            overdue=overdue, assets=[f"component {r.get('component', '')}"], workloads=[], scanners=[],
            detector=f"{TOOL_NAME} control assertion {aid}", fixable=True, vendor_product="", agreement=None,
            cvss=None, url="", plan=f"Change the {r.get('component', '')} configuration so that the assertion "
                                      f"passes ({r.get('title', aid)}), then re-run the control assertions.",
            mitigation="", identification_source=f"Control evidence engine - {TOOL_NAME} (run {run.get('id')})",
            comments=" ".join(filter(None, [
                f"[External UID SP-CTL-{aid}]", f"Evidence checked {r.get('checkedAt') or ''}.",
                f"SLA {v.sla_days.get(severity)}d from first failure {first_seen:%Y-%m-%d}.",
                "OVERDUE." if overdue else ""])),
            impact="",
        ))
    return items


def poam_baseline(v: View) -> str:
    return str(v.options.get("baseline") or v.engine.get("baseline") or v.engine_run.get("baseline")
               or "moderate").lower()


def _to_baseline(v: View, items: list[SimpleNamespace]) -> list[SimpleNamespace]:
    """eMASS only accepts POA&M items against controls in the system's control set (M4): keep the
    in-baseline controls (primary first), drop items left without one and count what was dropped."""
    baseline = poam_baseline(v)
    kept, dropped_items, dropped_tags = [], [], {}
    for i in items:
        inside = [c for c in i.controls if in_baseline(c, baseline)]
        outside = [c for c in i.controls if c not in inside]
        for c in outside:
            dropped_tags[c] = dropped_tags.get(c, 0) + 1
        if not inside:
            dropped_items.append(i)
            continue
        i.controls, i.controls_outside = inside, outside
        if outside:
            i.comments += f" Controls outside the {baseline} baseline not listed: {', '.join(outside)}."
        kept.append(i)
    v.poam_stats = {"baseline": baseline, "droppedItems": len(dropped_items),
                    "droppedItemIds": [i.poam_id for i in dropped_items][:50], "droppedTags": dropped_tags}
    return kept


def build_items(v: View) -> list[SimpleNamespace]:
    items = _vuln_items(v, granularity(v)) + _posture_items(v) + _stig_items(v) + _assertion_items(v)
    items = _to_baseline(v, items)
    items.sort(key=lambda i: (-sev_rank(i.severity), i.kind != "vulnerability", i.due or v.generated_at,
                              i.source_id))
    return items


# --------------------------------------------------------------------------- rows
def _assets_text(i: SimpleNamespace) -> str:
    txt = "\n".join(i.assets)
    if i.kind in ("vulnerability", "stig") and i.workloads:
        txt += "\nUsed by: " + ", ".join(i.workloads)
    return _cut(txt)


def _milestone(i: SimpleNamespace, v: View) -> str:
    """Milestone / Recommendations text: the fix (S1: Mitigations is for compensating measures in place)."""
    if i.kind == "vulnerability":
        return f"{i.plan} Target: {mdy(i.due)}."
    if i.kind == "assertion":
        return f"Correct the platform configuration and confirm by re-running the control assertions by {mdy(i.due)}."
    if i.kind == "stig":
        return f"Harden the image build for this STIG rule and confirm by re-evaluation (SCAP scan) by {mdy(i.due)}."
    return f"Remediate workload configuration and confirm by rescan by {mdy(i.due)}."


def _overdue_note(i: SimpleNamespace) -> str:
    """S1: an item whose scheduled date has passed needs a milestone change; flag it for the ISSO."""
    if not i.overdue:
        return ""
    return (f"OVERDUE: scheduled completion {mdy(i.due)} has passed. ISSO action: record a milestone change "
            "with the reason for the delay and a new date (many eMASS instances reject past dates on new items).")


def _resources(i: SimpleNamespace) -> str:
    return RESOURCES_DEFAULT if i.fixable else RESOURCES_VENDOR


def _tool_cols(i: SimpleNamespace, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if hasattr(x, "date") and x else x) if as_date else mdy
    return [i.poam_id, "Yes" if i.kev else "No", d(i.kev_due) if i.kev_due else "", IAVM_NOTE]


def emass_profile(v: View) -> dict[str, Any]:
    name = str(v.options.get("emassProfile") or "generic").lower()
    if name not in EMASS_PROFILES:
        from .registry import UnsupportedReport

        raise UnsupportedReport(f"unknown emassProfile {name!r} ({', '.join(EMASS_PROFILES)})")
    return EMASS_PROFILES[name]


def emass_columns(v: View) -> list[str]:
    prof = emass_profile(v)
    cols = [prof.get("rename", {}).get(c, c) for c in EMASS_COLUMNS if c not in prof.get("omit", set())]
    return cols + TOOL_COLUMNS


def emass_row(i: SimpleNamespace, v: View, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if x else None) if as_date else mdy
    poc = ", ".join(filter(None, [v.system.organization, v.system.poc_name, v.system.poc_email]))
    lvl = EMASS_LEVEL[i.severity]
    extra = f" Additional controls: {', '.join(i.controls[1:])}." if len(i.controls) > 1 else ""
    full = dict(zip(EMASS_COLUMNS, [
        "",                                   # POA&M Item ID - assigned by eMASS on import (External UID at the end)
        _cut(i.description),
        i.controls[0] if i.controls else "",  # one Control / AP per item: the primary in-baseline control
        _cut(i.security_checks, 30000),
        "Ongoing",
        d(i.due), "", "",
        1, _milestone(i, v), "Pending", _overdue_note(i), d(i.due), "",
        i.identification_source, _cut(i.detector), poc,
        _resources(i),
        _cut(i.comments + extra),
        lvl,                                  # Raw Severity (tool); Severity below is left for the ISSO
        _assets_text(i),
        "",                                   # Mitigations: compensating measures already in place (ISSO)
        "", "",                               # Predisposing Conditions, Severity (assessed, after mitigations)
        "", "", "", "",                       # Relevance of Threat, Threat Description, Likelihood, Impact
        "", "", _cut(_milestone(i, v)), "",   # Impact Description, Residual Risk Level, Recommendations, Resulting
    ]))
    prof = emass_profile(v)
    return [full[c] for c in EMASS_COLUMNS if c not in prof.get("omit", set())] + _tool_cols(i, as_date)


def generic_row(i: SimpleNamespace, v: View, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if x else None) if as_date else mdy
    poc = ", ".join(filter(None, [v.system.poc_name, v.system.poc_email])) or v.system.organization
    db_dates = [s.db_updated_at for s in v.scanners if s.name in i.scanners and s.db_updated_at]
    return [
        i.poam_id, ", ".join(i.controls), _cut(i.name), _cut(i.description), _cut(i.detector),
        _cut(i.security_checks, 30000), _assets_text(i), poc, _resources(i),
        _cut(i.plan), d(i.first_seen), d(i.due), _milestone(i, v), _overdue_note(i), d(v.generated_at),
        "No" if i.fixable else "Yes",
        d(max(db_dates)) if (db_dates and not i.fixable) else ("" if as_date else ""),
        i.vendor_product, FEDRAMP_RISK[i.severity], "", "No", "No", "No", "",
        f"{TOOL_NAME} {v.run_stamp} (assessment summary / vuln-export)", _cut(i.comments), "No",
        *_tool_cols(i, as_date),
    ]


def legacy_row(i: SimpleNamespace, v: View, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if x else None) if as_date else mdy
    org = ", ".join(filter(None, [v.system.organization, v.system.poc_name, v.system.poc_email]))
    return [
        _cut(i.description), ", ".join(i.controls), org, _cut(i.security_checks, 30000), _resources(i),
        d(i.due), f"1: {_milestone(i, v)}", _overdue_note(i), i.identification_source, "Ongoing", _cut(i.comments),
        CAT[i.severity], _assets_text(i), "", "", "",
        "", "", "", "", "", "", _cut(_milestone(i, v)), "",
    ]


VARIANTS = {
    "emass": ("eMASS", EMASS_COLUMNS, emass_row),
    "generic": ("POA&M", GENERIC_COLUMNS, generic_row),
    "emass-legacy": ("eMASS (legacy)", LEGACY_COLUMNS, legacy_row),
}


def _columns(v: View, variant: str) -> list[str]:
    return emass_columns(v) if variant == "emass" else VARIANTS[variant][1]


# --------------------------------------------------------------------------- writers
def _csv(v: View, items: list, variant: str) -> bytes:
    _sheet, _cols, rowf = VARIANTS[variant]
    cols = _columns(v, variant)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(cols)
    for i in items:
        w.writerow(safe_row(rowf(i, v, False)))
    return buf.getvalue().encode("utf-8-sig")


def _xlsx(v: View, items: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E7C4FF")
    overdue_fill = PatternFill("solid", fgColor="FDE2E1")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    date_cols = {"POA&M Scheduled Completion Date", "Milestone Scheduled Completion Date",
                 "Original Detection Date", "Scheduled Completion Date", "Status Date",
                 "Last Vendor Check-in Date", "KEV Due Date"}
    wide = {"Control Vulnerability Description", "Weakness Description", "Devices Affected", "Asset Identifier",
            "Comments", "Overall Remediation Plan", "Recommendations", "Milestone Description",
            "Mitigations (in-house and in conjunction with the Navy CSSP)", "Mitigations", "Impact Description"}

    scratch = wb.create_sheet("_styles")
    styles = {}
    for is_date in (False, True):
        for overdue in (False, True):
            c = scratch.cell(row=1 + len(styles), column=1)
            c.alignment = wrap
            if is_date:
                c.number_format = "mm/dd/yyyy"
            if overdue:
                c.fill = overdue_fill
            styles[(is_date, overdue)] = copy(c._style)
    wb.remove(scratch)

    for variant in ("emass", "generic", "emass-legacy"):
        title, _cols, rowf = VARIANTS[variant]
        cols = _columns(v, variant)
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = bold, head_fill, Alignment(wrap_text=True, vertical="center")
        # Performance (25k findings x 3 sheets): ws.max_row / ws[r] (via max_column) are O(cells),
        # and assigning style objects per cell hashes them against the workbook registry. Track
        # the row number, fetch cells with ws.cell (a dict lookup) and copy pre-registered
        # StyleArrays instead.
        for r, i in enumerate(items, start=2):
            ws.append(safe_row(rowf(i, v, True)))
            for idx, col in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=idx)
                cell._style = copy(styles[(col in date_cols and bool(cell.value), bool(i.overdue))])
        for idx, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = 60 if col in wide else max(14, min(30, len(col) + 2))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(items) + 1}"

    info = wb.create_sheet("Info")
    sc = {s: 0 for s in SEVERITIES}
    for i in items:
        sc[i.severity] += 1
    rows = [
        *([("WARNING", "Office/Org is empty: eMASS rejects POA&M items without it. Set settings `organization` "
                       "and regenerate before importing.")] if not (v.system.organization or "").strip() else []),
        ("Report", "Plan of Action & Milestones"),
        ("System / Project Name", v.system.name),
        ("Organization (Office/Org)", v.system.organization),
        ("eMASS System ID", v.system.emass_system_id),
        ("POC", ", ".join(filter(None, [v.system.poc_name, v.system.poc_email, v.system.poc_phone]))),
        ("Scope", v.scope_label),
        ("Scan ID", v.scan.id),
        ("Control evidence run", v.engine_run.get("id") or "none attached"),
        ("Generated from", v.run_stamp),
        ("Control set (baseline)", getattr(v, "poam_stats", {}).get("baseline", "")),
        ("Items dropped (no control in the baseline)", getattr(v, "poam_stats", {}).get("droppedItems", 0)),
        ("Control tags dropped (outside the baseline)",
         ", ".join(f"{c} x{n}" for c, n in sorted(getattr(v, "poam_stats", {}).get("droppedTags", {}).items()))
         or "none"),
        ("Scan finished", v.scan.finished_at.replace(tzinfo=None) if v.scan.finished_at else ""),
        ("Generated", v.generated_at.replace(tzinfo=None)),
        ("Generated by", TOOL_NAME),
        ("POA&M items", len(items)),
        ("Granularity", {"repository": "one item per image repository (remediation unit)",
                         "cve": "one item per vulnerability (Devices Affected lists the images)",
                         "finding": "one item per image x vulnerability x package"}[granularity(v)]),
        ("eMASS profile", str(v.options.get("emassProfile") or "generic")),
        ("Overdue items", sum(1 for i in items if i.overdue)),
        ("KEV items", sum(1 for i in items if i.kev)),
        ("KEV catalog", _kev_version()),
        *[(f"Items - {s}", n) for s, n in sc.items() if n],
        *[(f"SLA days - {s}", d) for s, d in v.sla_days.items()],
        ("", ""),
        ("Notes", "Machine-generated starting point for ISSO review; not a substitute for an assessor. "
                  "Sheet 'eMASS' follows the current eMASS POA&M import columns (header per emassProfile): paste "
                  "its template columns into the POA&M import template downloaded from your eMASS system. The "
                  "trailing columns External UID, KEV, KEV Due Date and IAVM ID are not template columns: drop "
                  "them when pasting, and use External UID to match items on later imports (or as the eMASS "
                  "REST API externalUid) so re-imports do not create duplicates. POA&M Item ID is left blank so "
                  "eMASS assigns IDs. Mitigations (compensating measures in place), Severity (assessed) and the "
                  "risk analysis columns are left for the ISSO; the fix is in Milestone Description and "
                  "Recommendations. Resources Required is a default to confirm. IAVM IDs are not available "
                  "from the scanners (column left blank). Overdue rows are shaded red and flagged in Milestone "
                  "Status Comments for a milestone change. KEV due dates (CISA BOD 22-01) override the severity "
                  "SLA where earlier."),
    ]
    for k, val in rows:
        info.append(safe_row([k, val]))
        info.cell(row=info.max_row, column=1).font = bold
        if isinstance(val, datetime):
            info.cell(row=info.max_row, column=2).number_format = "yyyy-mm-dd hh:mm"
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 100
    info.cell(row=info.max_row, column=2).alignment = wrap

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _kev_version() -> str:
    from .kev import catalog

    cat = catalog()
    return f"{cat.get('version') or 'unknown'} ({cat.get('source')})"


def generate(fmt: str, snapshot: Any, options: dict[str, Any]) -> GeneratedReport:
    v = normalize(snapshot, options)
    if not (v.system.organization or "").strip():
        # S1: eMASS requires Office/Org on every POA&M item: refuse on request, otherwise warn loudly
        if options.get("requireOrganization"):
            from .registry import UnsupportedReport

            raise UnsupportedReport("POA&M needs the Office/Org: set settings `organization`")
        from ..logs import get_logger

        get_logger(__name__).warning("poam.office_org_missing", hint="set settings organization")
    emass_profile(v)
    items = build_items(v)
    if fmt == "xlsx":
        return GeneratedReport(_xlsx(v, items), filename(v, "poam", "xlsx"), XLSX_CT)
    variant = str(options.get("poamVariant") or "emass").lower()
    if variant not in VARIANTS:
        from .registry import UnsupportedReport

        raise UnsupportedReport(f"unknown poamVariant {variant!r}")
    suffix = "" if variant == "emass" else f"-{variant}"
    return GeneratedReport(_csv(v, items, variant), filename(v, f"poam{suffix}", "csv"), CSV_CT)


def poam_controls(v: View) -> set[str]:
    """Controls that have at least one POA&M item in this view (SSP `planned` vs `not-implemented`)."""
    return {c for i in build_items(v) for c in i.controls}
